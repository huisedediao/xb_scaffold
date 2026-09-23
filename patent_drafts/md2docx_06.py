#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""将专利 06 的 Markdown 转为 Word(.docx)。宋体、小四、图片居中带图题。

与 01 模板的差异：
1. 图题行（"图N 标题"）自动在图题上方插入对应附图 figures_06/figN.png；
2. 权利要求等编号条目（"1. "）保留原编号，不替换为圆点。
"""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image as PILImage

BASE = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(BASE, "06_基于调试会话自动发现的界面源码跨进程定位方法及系统.md")
DOCX = os.path.join(BASE, "06_基于调试会话自动发现的界面源码跨进程定位方法及系统.docx")

SONG = "宋体"
TIMES = "Times New Roman"
SZ_BODY = Pt(12)      # 小四
SZ = {1: Pt(18), 2: Pt(15), 3: Pt(13)}

MAXW = Inches(6.3)
MAXH = Inches(7.8)

doc = Document()
style = doc.styles["Normal"]
style.font.name = TIMES
style.font.size = SZ_BODY
style.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)


def force_ea(p):
    """强制段落所有 run 及段落标记使用中文(宋体)字体。"""
    pr = p._p.get_or_add_pPr()
    rpr = pr.find(qn("w:rPr"))
    if rpr is None:
        rpr = pr.makeelement(qn("w:rPr"), {})
        pr.append(rpr)
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), SONG)
    for r in p.runs:
        rp = r._r.get_or_add_rPr()
        rrf = rp.find(qn("w:rFonts"))
        if rrf is None:
            rrf = rp.makeelement(qn("w:rFonts"), {})
            rp.append(rrf)
        rrf.set(qn("w:eastAsia"), SONG)


def add_run(p, text, size=SZ_BODY, bold=False):
    r = p.add_run(text)
    r.font.name = TIMES
    r.font.size = size
    r.font.bold = bold
    rp = r._r.get_or_add_rPr()
    rf = rp.find(qn("w:rFonts"))
    if rf is None:
        rf = rp.makeelement(qn("w:rFonts"), {})
        rp.append(rf)
    rf.set(qn("w:eastAsia"), SONG)
    return r


def add_inline(p, text, size=SZ_BODY):
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*", text):
        if m.start() > pos:
            add_run(p, text[pos:m.start()], size)
        add_run(p, m.group(1), size, bold=True)
        pos = m.end()
    if pos < len(text):
        add_run(p, text[pos:], size)


def fit_pic(sp, path):
    with PILImage.open(path) as im:
        w, h = im.size
    nw = MAXW
    nh = nw * h / w
    if nh > MAXH:
        nh = MAXH
        nw = nh * w / h
    sp.width = int(nw)
    sp.height = int(nh)


def add_figure(no, cap):
    path = os.path.join(BASE, "figures_06", "fig%d.png" % no)
    if not os.path.exists(path):
        print("WARN: figure not found:", path)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp = p.add_run().add_picture(path)
    fit_pic(sp, path)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(c, "图%d  %s" % (no, cap), size=Pt(10.5))
    force_ea(c)


lines = open(MD, encoding="utf-8").read().split("\n")
i = 0
n = len(lines)
while i < n:
    ln = lines[i]
    s = ln.strip()
    i += 1
    if not s:
        continue
    m = re.match(r"^(#{1,3})\s+(.*)", s)
    if m:
        lv = len(m.group(1))
        p = doc.add_paragraph()
        add_run(p, m.group(2).strip(), size=SZ.get(lv, SZ_BODY), bold=True)
        if lv == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        force_ea(p)
        continue
    if s.startswith("!["):
        mimg = re.match(r"!\[(.*?)\]\((.*?)\)", s)
        if mimg:
            alt, src = mimg.group(1), mimg.group(2)
            path = os.path.join(BASE, src)
            if os.path.exists(path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                sp = p.add_run().add_picture(path)
                fit_pic(sp, path)
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_run(cap, alt, size=Pt(10.5))
                force_ea(cap)
            else:
                print("WARN: image not found:", path)
        continue
    # 图题行：插入对应附图（"图N 为..."之类的说明行不在此列）
    mfig = re.match(r"^图(\d+)\s+([^为].*)$", s)
    if mfig and not s.startswith("摘要附图"):
        add_figure(int(mfig.group(1)), mfig.group(2).strip())
        continue
    if s.startswith("- "):
        p = doc.add_paragraph()
        add_run(p, "• ", SZ_BODY)
        add_inline(p, s[2:], SZ_BODY)
        p.paragraph_format.left_indent = Pt(18)
        force_ea(p)
        continue
    mnum = re.match(r"^(\d+)[\.、]\s+(.*)", s)
    if mnum:
        # 编号条目（权利要求等）：保留原编号
        p = doc.add_paragraph()
        add_run(p, mnum.group(1) + ". ", SZ_BODY)
        add_inline(p, mnum.group(2), SZ_BODY)
        p.paragraph_format.left_indent = Pt(18)
        force_ea(p)
        continue
    if s.startswith("|"):
        continue
    p = doc.add_paragraph()
    add_inline(p, s, SZ_BODY)
    p.paragraph_format.first_line_indent = Pt(24)
    force_ea(p)

doc.save(DOCX)
print("SAVED", DOCX)
